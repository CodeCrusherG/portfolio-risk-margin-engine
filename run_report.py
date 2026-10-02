"""Run the full pipeline: data -> SQL -> risk -> back-test -> margin -> Excel."""
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import risk_engine as re_

# Set TICKERS to real data when you run this on your own machine, e.g.
# ["RELIANCE.NS", "TCS.NS", "^NSEI", "USDINR=X", "LIQUIDBEES.NS"]
TICKERS = ["RELIANCE.NS", "TCS.NS", "^NSEI", "USDINR=X", "LIQUIDBEES.NS"]
WEIGHTS = {"RELIANCE.NS": 0.25, "TCS.NS": 0.20, "^NSEI": 0.30,
           "USDINR=X": 0.10, "LIQUIDBEES.NS": 0.15}
NOTIONAL = 10_000_000
ALPHA = 0.99


def main():
    prices = re_.load_prices(TICKERS)
    re_.save_to_sql(prices)
    prices = re_.query_prices()          # read back through SQL layer
    r = re_.portfolio_returns(prices, WEIGHTS)
    asset_rets = re_.np.log(prices / prices.shift(1)).dropna()

    mc_v, mc_e = re_.mc_var_es(asset_rets, WEIGHTS, ALPHA)
    summary = pd.DataFrame({
        "VaR 99% (1d)": [re_.hist_var(r, ALPHA), re_.param_var(r, ALPHA), mc_v],
        "ES 99% (1d)": [re_.hist_es(r, ALPHA), re_.param_es(r, ALPHA), mc_e],
    }, index=["Historical", "Parametric", "Monte Carlo"])
    summary["VaR (amount)"] = summary["VaR 99% (1d)"] * NOTIONAL
    summary["ES (amount)"] = summary["ES 99% (1d)"] * NOTIONAL

    bt = re_.backtest(r, 250, ALPHA)
    bt_df = pd.DataFrame({
        "Metric": ["Observations", "Exceptions", "Expected exceptions",
                   "Last-250d exceptions", "Basel zone", "Kupiec LR",
                   "Kupiec p-value", "Kupiec rejects model (5%)"],
        "Value": [bt["n_obs"], bt["n_exceptions"], bt["expected"],
                  bt["last250_exceptions"], bt["zone"],
                  round(bt["kupiec_lr"], 3), round(bt["kupiec_p"], 4),
                  bt["kupiec_reject"]]})

    exc_days = bt["exceptions"][bt["exceptions"]].index
    exc_df = pd.DataFrame({
        "Date": exc_days,
        "Return %": (r.loc[exc_days] * 100).round(2).values,
        "VaR %": (bt["var_series"].loc[exc_days] * 100).round(2).values})

    im = re_.margin_history(r, NOTIONAL)
    note = re_.commentary(r, im, NOTIONAL)
    stress = {"2008-style": {"RELIANCE.NS": -0.40, "TCS.NS": -0.40, "^NSEI": -0.35,
                         "USDINR=X": 0.15, "LIQUIDBEES.NS": 0.0},
          "Rate shock": {"LIQUIDBEES.NS": -0.01, "^NSEI": -0.05}}
    stress_df = pd.DataFrame(
        {"Scenario": list(stress),
         "P&L (amount)": [re_.stress_test(prices, WEIGHTS, s) * NOTIONAL
                          for s in stress.values()]})

    out = "risk_report.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.round(5).to_excel(xw, sheet_name="VaR_ES")
        bt_df.to_excel(xw, sheet_name="Backtest", index=False)
        exc_df.to_excel(xw, sheet_name="Exceptions", index=False)
        im.round(0).to_frame().to_excel(xw, sheet_name="InitialMargin")
        stress_df.to_excel(xw, sheet_name="Stress", index=False)
        pd.DataFrame({"Commentary": [note]}).to_excel(
            xw, sheet_name="Commentary", index=False)
    wb = load_workbook(out)
    for ws in wb:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 26
    wb["Commentary"].column_dimensions["A"].width = 140
    wb.save(out)

    fig, ax = plt.subplots(2, 1, figsize=(10, 7), sharex=False)
    ax[0].plot(r.loc[bt["var_series"].index], lw=0.6, label="Return")
    ax[0].plot(-bt["var_series"], color="red", lw=1, label="-VaR 99%")
    ax[0].scatter(exc_days, r.loc[exc_days], color="black", s=14,
                  label="Exception", zorder=3)
    ax[0].set_title(f"VaR back-test ({bt['zone']} zone, "
                    f"{bt['n_exceptions']} exceptions)")
    ax[0].legend()
    ax[1].plot(im, color="navy")
    ax[1].set_title("Initial margin (99%, 10-day), last 30 days")
    plt.tight_layout()
    plt.savefig("backtest_plot.png", dpi=130)

    print(summary.round(4), "\n")
    print(bt_df.to_string(index=False), "\n")
    print(stress_df, "\n")
    print(note)


if __name__ == "__main__":
    main()
